import io, json, sys, tempfile, unittest, zipfile
from pathlib import Path
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'workers/parser'))
import sandbox
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

class ParserTests(unittest.TestCase):
    def parse(self, data, ext):
        with tempfile.TemporaryDirectory() as temp:
            Path(temp, 'input').write_bytes(data)
            return json.loads(sandbox.run(temp, ext))
    def text(self, result): return '\n'.join(c['content'] for c in result['chunks'])
    def test_utf8_utf16_gb18030_and_html_are_plain_text(self):
        for encoding in ('utf-8', 'utf-16', 'gb18030'):
            r = self.parse('金额 123.45\n<script>alert(1)</script>'.encode(encoding), 'md')
            self.assertIn('金额 123.45', self.text(r)); self.assertIn('<script>', self.text(r))
            self.assertEqual('第 1–2 行', r['chunks'][0]['locator'])
    def test_long_markdown_keeps_all_lines_without_exhausting_fragment_limit(self):
        lines=[f'条款{i:04d}：原始业务说明。' for i in range(1100)]
        r=self.parse('\n'.join(lines).encode(),'md')
        self.assertEqual('READY',r['status']);self.assertEqual([],r['warnings'])
        self.assertEqual(lines,self.text(r).splitlines())
        self.assertLess(len(r['chunks']),500)
        self.assertTrue(all(len(c['content'])<=2000 for c in r['chunks']))
        self.assertTrue(r['chunks'][-1]['locator'].endswith('1100 行'))
    def test_csv_preserves_numeric_strings_and_multiline_cells(self):
        r=self.parse('名称,金额\n"两行\n备注",00123.4500\n'.encode(), 'csv')
        self.assertEqual('READY',r['status']);self.assertIn('00123.4500',self.text(r));self.assertEqual(2,len(r['chunks']))
    def test_blank_csv_has_no_fake_content_and_keeps_record_locations(self):
        r=self.parse(b'\n,\n" ",""\n', 'csv')
        self.assertEqual('FAILED',r['status']);self.assertIn('NO_TEXT_FOUND',r['warnings'])
        self.assertEqual([],r['chunks'])
        r=self.parse(b'\n,\namount,00123.4500\n', 'csv')
        self.assertEqual(1,len(r['chunks']));self.assertEqual('第 3 行',r['chunks'][0]['locator'])
    def test_docx_body_table_and_image_limitation(self):
        d=Document();d.add_paragraph('梦核资料基线');d.add_table(rows=1,cols=2).cell(0,0).text='105.20'
        d.sections[0].header.paragraphs[0].text='页眉未读取'
        from docx.oxml import OxmlElement
        d.paragraphs[0]._p.append(OxmlElement('w:ins'))
        out=io.BytesIO();d.save(out);r=self.parse(out.getvalue(),'docx')
        self.assertIn('梦核资料基线',self.text(r));self.assertIn('105.20',self.text(r));self.assertIn('OFFICE_NON_BODY_CONTENT',r['warnings'])
    def test_xlsx_values_and_uncalculated_formula_are_explicit(self):
        b=Workbook();s=b.active;s.title='月报';s.append(['金额',123.45]);s.append(['合计','=SUM(B1:B1)'])
        out=io.BytesIO();b.save(out);r=self.parse(out.getvalue(),'xlsx')
        self.assertEqual('PARTIAL',r['status']);self.assertIn('123.45',self.text(r));self.assertIn('FORMULA_VALUE_MISSING',r['warnings']);self.assertIn('月报',r['chunks'][0]['locator'])
    def xlsx_with_wrong_dimension(self, coordinate='B3'):
        import re
        b=Workbook();s=b.active;s['A1']='heading';s[coordinate]='dimension-marker-789'
        original=io.BytesIO();b.save(original);out=io.BytesIO()
        with zipfile.ZipFile(original) as src, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dest:
            for item in src.infolist():
                data=src.read(item.filename)
                if item.filename=='xl/worksheets/sheet1.xml':
                    data=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:A1"',data)
                dest.writestr(item,data)
        return out.getvalue()
    def test_xlsx_search_passage_keeps_title_headers_and_comparable_rows_together(self):
        b=Workbook();s=b.active;s.title='收款与订单'
        s.append(['月度收款同口径底表']);s.append(['统计期间','门店','收款（元）'])
        for month in ['2026-08','2026-09']:
            for store in range(5): s.append([month,f'合成门店{store}',1000+store])
        s.append([]);s.append(['另一张表']);s.append(['不能混入月度表',123])
        out=io.BytesIO();b.save(out);r=self.parse(out.getvalue(),'xlsx')
        self.assertEqual('READY',r['status'])
        matched=[c for c in r['chunks'] if '月度收款' in c['content']]
        self.assertEqual(1,len(matched));passage=matched[0]
        self.assertEqual('收款与订单 · 第 1–12 行',passage['locator'])
        self.assertIn('收款（元）',passage['content'])
        self.assertIn('A3="2026-08"',passage['content']);self.assertIn('A12="2026-09"',passage['content'])
        self.assertNotIn('另一张表',passage['content'])
    def test_xlsx_long_tables_keep_exact_cells_bounded_with_preceding_context(self):
        b=Workbook();s=b.active;s.title='测试'
        s.append(['收款表']);s.append(['门店','收款元'])
        for i in range(150): s.append([f'门店{i:03d}',1000+i])
        out=io.BytesIO();b.save(out);r=self.parse(out.getvalue(),'xlsx')
        self.assertEqual('READY',r['status']);self.assertGreater(len(r['chunks']),1)
        self.assertTrue(all(len(c['content'])<=2000 for c in r['chunks']))
        self.assertTrue(all('收款表' in c['content'] and '收款元' in c['content'] for c in r['chunks']))
        for i in range(150): self.assertEqual(1,self.text(r).count(f'B{i+3}={1000+i}'))
        self.assertIn('本段前文',r['chunks'][1]['content'])
    def test_xlsx_incorrect_dimensions_do_not_silently_drop_cells(self):
        r=self.parse(self.xlsx_with_wrong_dimension(),'xlsx')
        self.assertIn('B3="dimension-marker-789"',self.text(r))
        self.assertEqual('Sheet · 第 3 行',r['chunks'][-1]['locator'])
    def test_xlsx_actual_dimensions_still_enforce_limits(self):
        for coordinate in ('A20001','GS1'):
            r=self.parse(self.xlsx_with_wrong_dimension(coordinate),'xlsx')
            self.assertIn('RESOURCE_LIMIT',r['warnings'])
            self.assertNotIn('dimension-marker-789',self.text(r))
    def test_xlsx_nonstandard_relationship_target_cannot_bypass_grid_limit(self):
        original=io.BytesIO(self.xlsx_with_wrong_dimension('GS1'));out=io.BytesIO()
        with zipfile.ZipFile(original) as src, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dest:
            for item in src.infolist():
                data=src.read(item.filename);name=item.filename
                if name=='xl/worksheets/sheet1.xml': name='xl/custom/sheet.xml'
                if name=='xl/_rels/workbook.xml.rels':
                    data=data.replace(b'/xl/worksheets/sheet1.xml',b'/xl/custom/sheet.xml')
                dest.writestr(name,data)
        r=self.parse(out.getvalue(),'xlsx')
        self.assertEqual('FAILED',r['status']);self.assertIn('RESOURCE_LIMIT',r['warnings'])
    def image(self):
        im=Image.new('RGB',(1600,400),'white');d=ImageDraw.Draw(im);font=ImageFont.truetype(str(ROOT/'apps/api/src/main/resources/fonts/DreamcoreDocumentSansSC-Regular.ttf') if sys.platform=='linux' else '/System/Library/Fonts/Supplemental/Arial.ttf',65)
        d.text((70,100),'Dreamcore invoice 123.45',fill='black',font=font);return im
    def test_chinese_ocr(self):
        im=Image.new('RGB',(1600,400),'white');d=ImageDraw.Draw(im)
        font=ImageFont.truetype(str(ROOT/'apps/api/src/main/resources/fonts/DreamcoreDocumentSansSC-Regular.ttf') if sys.platform=='linux' else '/System/Library/Fonts/STHeiti Medium.ttc',70)
        d.text((70,100),'梦核工作空间 金额 123.45',fill='black',font=font)
        out=io.BytesIO();im.save(out,format='PNG');r=self.parse(out.getvalue(),'png')
        self.assertIn('工作空间',self.text(r).replace(' ',''));self.assertIn('123.45',self.text(r))
    def test_xls_cached_values(self):
        import xlwt
        b=xlwt.Workbook();s=b.add_sheet('旧月报');s.write(0,0,'金额');s.write(0,1,123.45)
        out=io.BytesIO();b.save(out);r=self.parse(out.getvalue(),'xls')
        self.assertEqual('PARTIAL',r['status']);self.assertIn('123.45',self.text(r));self.assertIn('XLS_CACHED_VALUES_ONLY',r['warnings'])
    def test_pdf_text_layer(self):
        stream=b'BT /F1 18 Tf 40 100 Td (Dreamcore budget 123.45) Tj ET'
        objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
        data=b'%PDF-1.4\n';offsets=[0]
        for i,obj in enumerate(objects,1):offsets.append(len(data));data+=f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n'
        at=len(data);data+=b'xref\n0 6\n0000000000 65535 f \n'+b''.join(f'{x:010d} 00000 n \n'.encode() for x in offsets[1:])+f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{at}\n%%EOF'.encode()
        r=self.parse(data,'pdf');self.assertIn('123.45',self.text(r));self.assertEqual('TEXT',r['chunks'][0]['method']);self.assertEqual('第 1 页',r['chunks'][0]['locator'])
    def test_image_and_scan_pdf_ocr(self):
        for ext in ('png','jpeg','pdf'):
            out=io.BytesIO();self.image().save(out,format=ext.upper());r=self.parse(out.getvalue(),ext)
            self.assertEqual('PARTIAL',r['status'],r);self.assertIn('123.45',self.text(r));self.assertIn('OCR_REVIEW_REQUIRED',r['warnings']);self.assertEqual('OCR',r['chunks'][0]['method'])
    def test_phone_photo_exif_orientation_is_applied_before_ocr(self):
        im=self.image().transpose(Image.Transpose.ROTATE_90)
        exif=Image.Exif();exif[274]=6
        out=io.BytesIO();im.save(out,format='JPEG',exif=exif)
        r=self.parse(out.getvalue(),'jpg')
        self.assertIn('123.45',self.text(r));self.assertEqual('PARTIAL',r['status'])
    def test_corrupt_unsupported_empty_and_limits(self):
        self.assertEqual('FAILED',self.parse(b'broken','pdf')['status'])
        self.assertEqual('FAILED',self.parse(b'','txt')['status'])
        self.assertEqual('UNSUPPORTED',self.parse(b'unknown','pptx')['status'])
        r=self.parse((('中'*2000+'\n')*210).encode(),'txt');self.assertEqual('PARTIAL',r['status']);self.assertIn('CONTENT_LIMIT',r['warnings']);self.assertLessEqual(sum(len(c['content']) for c in r['chunks']),400_000)
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:z.writestr('word/document.xml',b'A'*1_000_000)
        self.assertEqual('FAILED',self.parse(out.getvalue(),'docx')['status'])
    def test_sandbox_denies_other_task_files_writes_and_network(self):
        with tempfile.TemporaryDirectory() as outside, tempfile.TemporaryDirectory() as task:
            canary=Path(outside,'secret');canary.write_text('synthetic canary')
            probe=ROOT/'workers/parser/.local/probe.py'
            probe.parent.mkdir(exist_ok=True)
            probe.write_text('import socket,json\nfrom pathlib import Path\nresult=[]\n'+
                f'for action in [lambda: Path({str(canary)!r}).read_text(), lambda: Path({str(canary)!r}).write_text("bad"), lambda: socket.create_connection(("127.0.0.1", 80), timeout=1)]:\n'+
                ' try: action();result.append(False)\n except OSError: result.append(True)\nprint(json.dumps(result))\n')
            try:self.assertEqual([True,True,True],json.loads(sandbox.run(task,'',probe)))
            finally:probe.unlink()
            self.assertEqual('synthetic canary',canary.read_text())

if __name__=='__main__':unittest.main()
