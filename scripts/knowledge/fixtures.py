"""Generate synthetic format-acceptance files; no customer files or network access."""
import hashlib, io, json, re, sys, zipfile
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
import xlwt

def saved(obj, **kwargs):
    out = io.BytesIO(); obj.save(out, **kwargs); return out.getvalue()

def image(chinese=False):
    im = Image.new('RGB', (1600, 400), 'white')
    font = '/System/Library/Fonts/STHeiti Medium.ttc' if chinese else '/System/Library/Fonts/Supplemental/Arial.ttf'
    ImageDraw.Draw(im).text((70, 100), '工作空间 金额 123.45' if chinese else 'Dreamcore invoice 123.45',
                           fill='black', font=ImageFont.truetype(font, 65))
    return im

def text_pdf():
    stream = b'BT /F1 18 Tf 40 100 Td (FORMATPDF invoice 123.45) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    data = b'%PDF-1.4\n'; offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data)); data += f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n'
    at = len(data)
    return data + b'xref\n0 6\n0000000000 65535 f \n' + b''.join(f'{x:010d} 00000 n \n'.encode() for x in offsets) + f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{at}\n%%EOF'.encode()

def stale_xlsx(coordinate='B3'):
    book = Workbook(); sheet = book.active; sheet['A1'] = '标题'; sheet[coordinate] = 'FORMATDIMENSION'
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(saved(book))) as src, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dest:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', data)
            dest.writestr(item, data)
    return out.getvalue()

def generate(directory):
    directory.mkdir(parents=True, exist_ok=True); manifest = []
    def case(name, data, status, warnings=(), query=None, locator=None, method=None):
        (directory / name).write_bytes(data)
        manifest.append(dict(name=name, sha256=hashlib.sha256(data).hexdigest(), status=status,
                             warnings=list(warnings), query=query, locator=locator, method=method))
    case('01-中文.TXT', 'FORMATTXT 梦核资料样例\n金额 00123.4500'.encode(), 'READY', query='FORMATTXT', locator='第 1 行', method='TEXT')
    case('02-文稿.md', '# FORMATMD\n<script>window.__formatInjected=true</script>'.encode(), 'READY', query='FORMATMD', locator='第 1 行', method='TEXT')
    case('03-UTF16.txt', 'FORMATUTF16 中文样例'.encode('utf-16'), 'READY', query='FORMATUTF16', locator='第 1 行', method='TEXT')
    case('04-GB18030.txt', 'FORMATGB 中文样例'.encode('gb18030'), 'PARTIAL', ['ENCODING_GB18030'], 'FORMATGB', '第 1 行', 'TEXT')
    case('05-多行及空行.csv', '\n,\n项目,金额\n"FORMATCSV\n多行备注",00123.4500\n'.encode(), 'READY', query='FORMATCSV', locator='第 4 行', method='TABLE')
    doc = Document(); doc.add_paragraph('FORMATDOCX 中文会议记录')
    doc.add_table(rows=1, cols=2).cell(0, 0).text = '00123.4500'
    doc.sections[0].header.paragraphs[0].text = '未完整提取的页眉'
    case('06-正文表格页眉.docx', saved(doc), 'PARTIAL', ['OFFICE_NON_BODY_CONTENT'], 'FORMATDOCX', '正文第 1 段', 'TEXT')
    book = Workbook(); sheet = book.active; sheet.title = '月报'
    sheet.append(['FORMATXLSX', 123.45]); sheet.append(['合计', '=SUM(B1:B1)'])
    case('07-公式无缓存.xlsx', saved(book), 'PARTIAL', ['FORMULAS_NOT_RECALCULATED', 'FORMULA_VALUE_MISSING'], 'FORMATXLSX', '月报 · 第 1 行', 'TABLE')
    book = xlwt.Workbook(); sheet = book.add_sheet('旧月报'); sheet.write(0, 0, 'FORMATXLS'); sheet.write(0, 1, 123.45)
    case('08-旧表格.xls', saved(book), 'PARTIAL', ['XLS_CACHED_VALUES_ONLY'], 'FORMATXLS', '旧月报 · 第 1 行', 'TABLE')
    case('09-文字层.pdf', text_pdf(), 'PARTIAL', ['PDF_LAYOUT_NOT_INTERPRETED'], 'FORMATPDF', '第 1 页', 'TEXT')
    case('10-扫描件.pdf', saved(image(), format='PDF'), 'PARTIAL', ['OCR_REVIEW_REQUIRED', 'PDF_LAYOUT_NOT_INTERPRETED'], '123.45', '第 1 页', 'OCR')
    case('11-中文图片.png', saved(image(True), format='PNG'), 'PARTIAL', ['OCR_REVIEW_REQUIRED'], '123.45', '图片文字', 'OCR')
    case('12-图片.jpeg', saved(image(), format='JPEG'), 'PARTIAL', ['OCR_REVIEW_REQUIRED'], '123.45', '图片文字', 'OCR')
    exif = Image.Exif(); exif[274] = 6
    case('13-旋转手机照片.jpg', saved(image().transpose(Image.Transpose.ROTATE_90), format='JPEG', exif=exif), 'PARTIAL', ['OCR_REVIEW_REQUIRED'], '123.45', '图片文字', 'OCR')
    case('14-范围标记错误.xlsx', stale_xlsx(), 'READY', query='FORMATDIMENSION', locator='Sheet · 第 3 行', method='TABLE')
    case('15-全空表格.csv', b'\n,\n" ",""\n', 'FAILED', ['NO_TEXT_FOUND'])
    case('16-损坏.pdf', b'synthetic corrupt PDF', 'FAILED', ['PARSE_FAILED'])
    case('17-损坏.docx', b'synthetic corrupt DOCX', 'FAILED', ['PARSE_FAILED'])
    case('18-不支持.pptx', b'synthetic unsupported file; not a usable presentation', 'UNSUPPORTED')
    case('19-内容超限.md', ('FORMATLIMIT\n' * 501).encode(), 'PARTIAL', ['CONTENT_LIMIT', 'RESOURCE_LIMIT'], 'FORMATLIMIT', '第 1 行', 'TEXT')
    case('20-表格超限.xlsx', stale_xlsx('GS1'), 'FAILED', ['RESOURCE_LIMIT'])
    case('21-空文件.txt', b'', 'FAILED', ['NO_TEXT_FOUND'])
    (directory / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')

if __name__ == '__main__': generate(Path(sys.argv[1]))
