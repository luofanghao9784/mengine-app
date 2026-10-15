import importlib.util, tempfile, unittest
from pathlib import Path
spec=importlib.util.spec_from_file_location('structured_table',Path(__file__).with_name('structured_table.py'))
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)

class Tables(unittest.TestCase):
 def parse(self,text,ext='csv'):
  with tempfile.TemporaryDirectory() as root:
   p=Path(root)/'input';p.write_bytes(text if isinstance(text,bytes) else text.encode());return m.extract(p,ext)
 def test_exact_csv(self):
  s=self.parse('\ufeff编号,金额,备注\r\n00123,0.1000,"第一行\n第二行,有逗号"\r\n00123,0.1000,\r\n')
  self.assertEqual('READY',s['state']);self.assertEqual(['text','number','text'],s['types']);self.assertEqual('00123',s['rows'][0][0]);self.assertEqual('0.1000',s['rows'][0][1]);self.assertEqual('',s['rows'][1][2]);self.assertEqual(2,len(s['rows']))
 def test_header_and_shape(self):
  for text in ('a,a\n1,2','a,\n1,2','a,b\n1','a\n"unterminated','a,b\na,b','a,b\n合计,1'):
   self.assertEqual('DOCUMENT_ONLY',self.parse(text)['state'],text)
 def test_no_silent_truncation(self):
  self.assertEqual('READY',self.parse('a\n'+'1\n'*10000)['state'])
  self.assertEqual('FILE_TABLE_LIMIT',self.parse('a\n'+'1\n'*10001)['reason'])
  self.assertEqual('DOCUMENT_ONLY',self.parse('a\n'+'x'*2001)['state'])
 def test_encoding_types_and_tsv(self):
  self.assertEqual('FILE_TABLE_ENCODING',self.parse('编号\n一'.encode('gb18030'))['reason'])
  self.assertEqual(['text','date','boolean','text'],self.parse('id\t日期\t开关\t金额\n001\t2026-01-01\ttrue\t0001','tsv')['types'])
  self.assertEqual('FILE_TABLE_LEGACY_XLS',self.parse(b'anything','xls')['reason'])
 def xlsx(self,mutate):
  import openpyxl
  with tempfile.TemporaryDirectory() as root:
   p=Path(root)/'input.xlsx';w=openpyxl.Workbook();s=w.active;s.append(['编号','金额']);s.append(['001',12.34]);mutate(w);w.save(p);w.close();return m.extract(p,'xlsx')
 def test_excel(self):
  self.assertEqual('READY',self.xlsx(lambda w:None)['state'])
  self.assertEqual('FILE_TABLE_NUMBER_PRECISION',self.xlsx(lambda w:setattr(w.active['B2'],'value',12345678901234.12))['reason'])
  for mutate in (lambda w:w.create_sheet(),lambda w:setattr(w.active['B2'],'value','=1+2'),lambda w:w.active.merge_cells('A3:B3'),lambda w:setattr(w.active['B2'],'number_format','0%'),lambda w:setattr(w.active.row_dimensions[2],'hidden',True)):
   self.assertEqual('DOCUMENT_ONLY',self.xlsx(mutate)['state'])

if __name__=='__main__':unittest.main()
