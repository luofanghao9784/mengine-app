"""Local parser protocol 1. All input is untrusted; launched only by sandbox.py."""
import resource
resource.setrlimit(resource.RLIMIT_CPU, (45, 45))
resource.setrlimit(resource.RLIMIT_FSIZE, (8_388_608, 8_388_608))
resource.setrlimit(resource.RLIMIT_NOFILE, (128, 128))
# macOS does not enforce RLIMIT_AS reliably; input/expansion/pixel limits remain mandatory.
import csv, io, json, os, re, subprocess, sys, zipfile, warnings
from pathlib import Path

VERSION = 'local-3'
chunks, notes = [], set()
MAX_CHARS, MAX_CHUNKS = 400_000, 500
used = 0
LOCAL_TOOLS = Path(__file__).resolve().parent / '.local/tools/bin'
TOOL_ROOT = '/usr/bin' if sys.platform == 'linux' else str(LOCAL_TOOLS) if LOCAL_TOOLS.is_dir() else '/opt/homebrew/bin'
class Limit(Exception): pass

def add(locator, text, method='TEXT'):
    global used
    text = str(text).replace('\x00', '').strip()
    for start in range(0, len(text), 2000):
        part = text[start:start + 2000]
        if len(chunks) >= MAX_CHUNKS or used + len(part) > MAX_CHARS:
            notes.add('CONTENT_LIMIT'); raise Limit()
        chunks.append({'locator': locator[:180], 'content': part, 'method': method})
        used += len(part)

def command(args):
    return subprocess.run(args, check=True, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=25).stdout.decode('utf-8', errors='strict')

def decode(path):
    data = path.read_bytes()
    if data.startswith((b'\xff\xfe', b'\xfe\xff')): return data.decode('utf-16')
    try: return data.decode('utf-8-sig')
    except UnicodeError:
        notes.add('ENCODING_GB18030'); return data.decode('gb18030')

def zip_check(path):
    with zipfile.ZipFile(path) as archive:
        files = archive.infolist()
        if len(files) > 3000 or sum(x.file_size for x in files) > 40_000_000:
            raise Limit()
        for x in files:
            if x.flag_bits & 1 or x.file_size > 12_000_000 or x.file_size > max(100_000, x.compress_size * 200): raise Limit()
            # Scan XML for entity declarations before parsers handle them; never extract the archive.
            if x.filename.endswith(('.xml', '.rels')):
                body = archive.read(x)
                if b'<!DOCTYPE' in body.upper() or b'<!ENTITY' in body.upper(): raise ValueError()
        return [x.filename for x in files]

def ocr(path, locator):
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = 20_000_000
    warnings.simplefilter('error', Image.DecompressionBombWarning)
    with Image.open(path) as im:
        if im.width * im.height > 20_000_000: raise Limit()
        im.verify()
    # Phone photos can store upright display orientation separately from pixel data.
    # Only the task-local copy is normalized; the uploaded original stays untouched.
    normalized = path.parent / 'ocr-oriented.png'
    try:
        with Image.open(path) as im:
            oriented = im.getexif().get(274, 1) in range(2, 9)
            if oriented: ImageOps.exif_transpose(im).convert('RGB').save(normalized)
        text = command([TOOL_ROOT + '/tesseract', str(normalized if oriented else path), 'stdout', '-l', 'chi_sim', '--psm', '3'])
    finally: normalized.unlink(missing_ok=True)
    notes.add('OCR_REVIEW_REQUIRED')
    add(locator, text, 'OCR')

def xlsx_grid_check(path, worksheets):
    # Check actual coordinates before openpyxl allocates rows. Declared dimensions
    # are often stale and cannot be trusted as either a content or resource bound.
    from defusedxml.ElementTree import iterparse
    with zipfile.ZipFile(path) as archive:
        for name in worksheets:
            rows, row_number, cells = 0, 0, 0
            with archive.open(name) as stream:
                for event, element in iterparse(stream, events=('start', 'end')):
                    tag = element.tag.rsplit('}', 1)[-1]
                    if event == 'start' and tag == 'row':
                        rows += 1; row_number = int(element.get('r', row_number + 1)); cells = 0
                        if rows > 20_000 or not 1 <= row_number <= 20_000: raise Limit()
                    if event == 'start' and tag == 'c':
                        cells += 1
                        ref = element.get('r')
                        if ref:
                            match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})', ref)
                            if not match: raise Limit()
                            column = 0
                            for char in match[1]: column = column * 26 + ord(char) - 64
                            if column > 200 or int(match[2]) > 20_000: raise Limit()
                        if cells > 200: raise Limit()
                    if event == 'end': element.clear()

class TablePassages:
    """Keep adjacent spreadsheet cells together, without inferring headers or values."""
    def __init__(self, title):
        self.title = title
        self.rows = []
        self.opening = []

    def flush(self):
        if not self.rows: return
        first, last = self.rows[0][0], self.rows[-1][0]
        locator = f'{self.title} · 第 {first} 行' if first == last else f'{self.title} · 第 {first}–{last} 行'
        context = [row for row in self.opening if row[0] < first]
        text = '\n'.join([f'工作表：{self.title}'] +
                         ([f'本段前文（第 {context[0][0]}–{context[-1][0]} 行）：'] + [r[1] for r in context] if context else []) +
                         [r[1] for r in self.rows])
        add(locator, text, 'TABLE')
        self.rows = []

    def append(self, number, text):
        if not text:
            self.flush(); self.opening = []; return
        # Preserve oversized rows through the existing bounded splitter. Do not
        # mistake a split cell for a complete table or discard the remaining bytes.
        if len(text) > 1200:
            self.flush(); self.opening = []
            add(f'{self.title} · 第 {number} 行', text, 'TABLE')
            return
        candidate = self.rows + [(number, text)]
        context = [r for r in self.opening if r[0] < candidate[0][0]]
        size = len(self.title) + 80 + sum(len(r[1]) + 1 for r in context + candidate)
        if size > 1950:
            self.flush()
        self.rows.append((number, text))
        # Repeat only exact, bounded opening rows on later passages. They are
        # labeled as preceding text, never promoted to inferred column headers.
        if len(self.opening) < 2 and sum(len(r[1]) for r in self.opening) + len(text) <= 600:
            self.opening.append((number, text))

def parse(path, extension):
    if extension in ('txt', 'md'):
        for i, line in enumerate(decode(path).splitlines(), 1): add(f'第 {i} 行', line)
    elif extension in ('csv', 'tsv'):
        for i, row in enumerate(csv.reader(io.StringIO(decode(path)), delimiter='\t' if extension == 'tsv' else ','), 1):
            if any(cell.strip() for cell in row):
                add(f'第 {i} 行', json.dumps(row, ensure_ascii=False), 'TABLE')
    elif extension == 'docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        files = zip_check(path)
        if any('/media/' in f or any(k in f for k in ('footnotes', 'endnotes', 'header', 'footer', 'comments')) for f in files):
            notes.add('OFFICE_NON_BODY_CONTENT')
        doc = Document(path)
        if doc.element.xpath('//w:ins | //w:del | //w:fldChar | //w:altChunk | //w:object | //w:txbxContent | //w:tbl//w:tbl'):
            notes.add('OFFICE_NON_BODY_CONTENT')
        p, t = 0, 0
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                p += 1; add(f'正文第 {p} 段', block.text)
                if block._element.xpath('.//w:txbxContent'): notes.add('OFFICE_NON_BODY_CONTENT')
            elif isinstance(block, Table):
                t += 1
                for r, row in enumerate(block.rows, 1): add(f'表 {t} · 第 {r} 行', json.dumps([c.text for c in row.cells], ensure_ascii=False), 'TABLE')
    elif extension == 'xlsx':
        import mimetypes
        mimetypes.knownfiles = []
        from openpyxl import load_workbook
        files = zip_check(path)
        if any('/drawings/' in f or '/charts/' in f or '/externalLinks/' in f for f in files): notes.add('OFFICE_NON_BODY_CONTENT')
        values = load_workbook(path.open('rb'), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(path.open('rb'), read_only=True, data_only=False, keep_links=False)
        try:
            # Pinned openpyxl 3.1.5 resolves relationship targets here. Inspect the
            # actual parts, including nonstandard paths, before iterating cells.
            xlsx_grid_check(path, [sheet._worksheet_path for sheet in formulas])
            for sheet in formulas:
                sheet.reset_dimensions(); values[sheet.title].reset_dimensions()
                passages = TablePassages(sheet.title)
                for i, (raw, cached) in enumerate(zip(sheet.iter_rows(), values[sheet.title].iter_rows()), 1):
                    if i > 20_000 or len(raw) > 200 or len(cached) > 200: raise Limit()
                    cells = []
                    for cell, value in zip(raw, cached):
                        v = value.value
                        if cell.data_type == 'f':
                            notes.add('FORMULAS_NOT_RECALCULATED')
                            v = {'formula': cell.value, 'cachedValue': v}
                            if value.value is None: notes.add('FORMULA_VALUE_MISSING')
                        if v is not None: cells.append(f'{cell.coordinate}={json.dumps(v, ensure_ascii=False, default=str)}')
                    passages.append(i, ' | '.join(cells))
                passages.flush()
        finally: values.close(); formulas.close()
    elif extension == 'xls':
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        notes.add('XLS_CACHED_VALUES_ONLY')
        try:
            for sheet in book.sheets():
                if sheet.nrows > 20_000 or sheet.ncols > 200: raise Limit()
                for i in range(sheet.nrows): add(f'{sheet.name} · 第 {i + 1} 行', json.dumps(sheet.row_values(i), ensure_ascii=False), 'TABLE')
        finally: book.release_resources()
    elif extension == 'pdf':
        info = command([TOOL_ROOT + '/pdfinfo', str(path)])
        count = int(re.search(r'^Pages:\s+(\d+)', info, re.M).group(1))
        if count > 30: notes.add('PAGE_LIMIT')
        # PDF text layers can omit diagrams or images even when some text exists.
        notes.add('PDF_LAYOUT_NOT_INTERPRETED')
        for page in range(1, min(count, 30) + 1):
            text = command([TOOL_ROOT + '/pdftotext', '-f', str(page), '-l', str(page), '-enc', 'UTF-8', str(path), '-'])
            if text.strip(): add(f'第 {page} 页', text)
            else:
                prefix = str(path.parent / 'page')
                command([TOOL_ROOT + '/pdftoppm', '-f', str(page), '-l', str(page), '-singlefile', '-scale-to', '1800', '-png', str(path), prefix])
                ocr(Path(prefix + '.png'), f'第 {page} 页')
                Path(prefix + '.png').unlink(missing_ok=True)
    elif extension in ('png', 'jpg', 'jpeg'): ocr(path, '图片文字')
    else: return 'UNSUPPORTED'
    return 'PARTIAL' if notes else 'READY'

if __name__ == '__main__':
    status = 'FAILED'
    try:
        path = Path(sys.argv[1]); extension = sys.argv[2]
        if path.stat().st_size > 20_971_520: raise Limit()
        status = parse(path, extension)
        if status != 'UNSUPPORTED' and not chunks:
            status = 'FAILED'; notes.add('NO_TEXT_FOUND')
    except Limit:
        notes.add('RESOURCE_LIMIT'); status = 'PARTIAL' if chunks else 'FAILED'
    except Exception:
        notes.add('PARSE_FAILED'); chunks.clear(); status = 'FAILED'
    import importlib.util
    spec = importlib.util.spec_from_file_location('structured_table', Path(__file__).with_name('structured_table.py'))
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    structured = module.extract(path, extension)
    print(json.dumps({'structured': structured, 'parserVersion': VERSION, 'status': status, 'warnings': sorted(notes), 'chunks': chunks}, ensure_ascii=False))
